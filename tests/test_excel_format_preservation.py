"""
Tests para la preservación de formato al sincronizar el Excel base.

Bug original: cada sync recargaba la salida anterior de openpyxl (lossy), por lo que
el formato del Excel (incl. los desplegables de validación) se degradaba en syncs
sucesivos. Fix: regenerar siempre desde una plantilla pristine guardada aparte.

Estos tests usan la plantilla REAL del cliente (Ekol Fase 2) y mockean la capa de BD
con un store en memoria.
"""

import os
import sys
import zipfile
from io import BytesIO

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import database
import excel_base


TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    "docs", "Ekol Fase 2", "62999_4500096466_Draft Mailer Template Newer 2.xlsx",
)

pytestmark = pytest.mark.skipif(
    not os.path.exists(TEMPLATE_PATH),
    reason="plantilla real del cliente no disponible",
)


def _count_validations(raw: bytes) -> int:
    z = zipfile.ZipFile(BytesIO(raw))
    xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8", "replace")
    return xml.count("<dataValidation ")


@pytest.fixture
def mock_db(monkeypatch):
    """Store en memoria que imita project_excel_files para un único proyecto."""
    store = {"template": None, "working": None, "filename": ""}

    def get_project_id(name):
        return 1

    def save_project_excel(project_id, excel_data, filename="", set_template=False):
        store["working"] = bytes(excel_data)
        if set_template:
            store["template"] = bytes(excel_data)
        if filename:
            store["filename"] = filename

    def load_project_excel(project_id):
        return store["working"]

    def load_project_excel_template(project_id):
        return store["template"] if store["template"] is not None else store["working"]

    monkeypatch.setattr(database, "get_project_id", get_project_id)
    monkeypatch.setattr(database, "save_project_excel", save_project_excel)
    monkeypatch.setattr(database, "load_project_excel", load_project_excel)
    monkeypatch.setattr(database, "load_project_excel_template", load_project_excel_template)
    return store


@pytest.fixture
def db_lines():
    return [
        {"paquete_num": 1, "marca": "S355-J0", "piezas": 100},
        {"paquete_num": 2, "marca": "S355-J2", "piezas": 50},
        {"paquete_num": 3, "marca": "P62989",  "piezas": 75},
    ]


def _upload(store):
    """Simula la subida del usuario: guarda plantilla + copia de trabajo."""
    raw = open(TEMPLATE_PATH, "rb").read()
    database.save_project_excel(1, raw, "template.xlsx", set_template=True)
    return raw


def test_validations_survive_first_and_second_sync(mock_db, db_lines):
    """El núcleo del bug: las 10 validaciones (desplegables) deben sobrevivir a CADA sync."""
    orig = _upload(mock_db)
    assert _count_validations(orig) == 10

    excel_base.write_lines_to_excel("PROY", db_lines)
    after_sync1 = mock_db["working"]
    assert _count_validations(after_sync1) == 10, "Sync 1 perdió validaciones"

    excel_base.write_lines_to_excel("PROY", db_lines)
    after_sync2 = mock_db["working"]
    assert _count_validations(after_sync2) == 10, "Sync 2 perdió validaciones (regresión del bug)"


def test_template_never_mutated_by_sync(mock_db, db_lines):
    """La plantilla pristine no debe cambiar nunca tras sincronizar."""
    _upload(mock_db)
    template_before = mock_db["template"]
    excel_base.write_lines_to_excel("PROY", db_lines)
    excel_base.write_lines_to_excel("PROY", db_lines)
    assert mock_db["template"] == template_before, "La plantilla fue sobrescrita por el sync"


def test_data_written_to_v2_columns(mock_db, db_lines):
    """Los datos se escriben en las columnas correctas del layout v2 (B/C/I/J)."""
    _upload(mock_db)
    excel_base.write_lines_to_excel("PROY", db_lines)
    wb = load_workbook(BytesIO(mock_db["working"]))
    ws = wb.active
    # Fila 14 = primera línea (paquete 1)
    assert ws.cell(14, 2).value == "4500096466/0001"   # B: paquete_code
    assert ws.cell(14, 3).value == "Bundle"            # C: tipo
    assert ws.cell(14, 10).value == 100                # J: piezas
    # 3 filas de datos en total, sin duplicados
    data_rows = sum(1 for r in range(14, ws.max_row + 1) if ws.cell(r, 2).value)
    assert data_rows == 3
    wb.close()


def test_resync_is_idempotent(mock_db, db_lines):
    """Sincronizar dos veces con las mismas líneas NO duplica filas (reflejo puro de la BD)."""
    _upload(mock_db)
    excel_base.write_lines_to_excel("PROY", db_lines)
    excel_base.write_lines_to_excel("PROY", db_lines)
    wb = load_workbook(BytesIO(mock_db["working"]))
    ws = wb.active
    data_rows = sum(1 for r in range(14, ws.max_row + 1) if ws.cell(r, 2).value)
    wb.close()
    assert data_rows == 3, f"Re-sync duplicó filas: {data_rows} (esperado 3)"


def test_deletion_reflected_after_resync(mock_db, db_lines):
    """Si se elimina una línea en BD, el siguiente sync la quita del Excel."""
    _upload(mock_db)
    excel_base.write_lines_to_excel("PROY", db_lines)
    # Eliminar la línea del paquete 2
    remaining = [l for l in db_lines if l["paquete_num"] != 2]
    excel_base.write_lines_to_excel("PROY", remaining)
    wb = load_workbook(BytesIO(mock_db["working"]))
    ws = wb.active
    codes = [ws.cell(r, 2).value for r in range(14, ws.max_row + 1) if ws.cell(r, 2).value]
    wb.close()
    assert "4500096466/0002" not in codes
    assert len(codes) == 2
